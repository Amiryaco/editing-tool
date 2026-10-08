#!/usr/bin/env python3
"""Analyse normalised transactions (from parse_statements.py): categories, business vs personal,
fixed charges and subscriptions, future installment commitments, anomalies, and a baseline budget.

    python3 analyze.py finance/out/transactions.csv --config finance/config.json --out finance/out

Writes <out>/analysis.json (everything the assistant needs to write the plan; every number in it
comes from the statements), <out>/categorized.csv and <out>/report.xlsx. report.py turns
analysis.json into the HTML dashboard.
"""
import argparse
import datetime as dt
import json
import os
import sys
from collections import defaultdict

import pandas as pd

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from categories import CATEGORIES, categorize, merchant_key  # noqa: E402

VAT_RATE = 0.18           # Israel, since 1 Jan 2025
UNCAT = "לא מסווג"


def load_config(path):
    cfg = {"cards": {}, "merchant_rules": [], "ignore": [], "budgets": {}, "income": {},
           "business": {}}
    if path and os.path.exists(path):
        with open(path, encoding="utf-8") as fh:
            cfg.update(json.load(fh))
    return cfg


def months_between(a, b):
    ya, ma = map(int, a.split("-"))
    yb, mb = map(int, b.split("-"))
    return (yb - ya) * 12 + (mb - ma)


def add_months(m, k):
    y, mo = map(int, m.split("-"))
    t = y * 12 + (mo - 1) + k
    return f"{t // 12:04d}-{t % 12 + 1:02d}"


def r0(x):
    return int(round(float(x)))


def prepare(df, cfg):
    df = df.copy()
    for c in ("merchant", "issuer_category", "card", "txn_type", "notes", "orig_currency", "kind"):
        df[c] = df[c].fillna("").astype(str)
    df["card"] = df["card"].str.replace(r"\.0$", "", regex=True).str.zfill(4).where(df["card"] != "", "")
    df["amount"] = pd.to_numeric(df["amount"], errors="coerce").fillna(0.0)
    df["orig_amount"] = pd.to_numeric(df["orig_amount"], errors="coerce")
    for c in ("installment_no", "installment_total"):
        df[c] = pd.to_numeric(df[c], errors="coerce")
    df["mkey"] = df["merchant"].map(merchant_key)
    cats = [categorize(m, ic, cfg["merchant_rules"]) for m, ic in zip(df["merchant"], df["issuer_category"])]
    df["category"] = [c[0] for c in cats]
    df["biz_hint"] = [c[1] for c in cats]
    df["cat_source"] = [c[2] for c in cats]
    df["nature"] = df["category"].map(lambda c: CATEGORIES.get(c, "lifestyle"))

    # scope: user rule > card mapping > suggested by merchant type > personal
    rule_scope = {}
    for r in cfg["merchant_rules"]:
        if r.get("scope"):
            rule_scope[r["match"]] = r["scope"]

    def scope(row):
        for match, sc in rule_scope.items():
            if match.upper() in row["merchant"].upper():
                return sc
        card = cfg["cards"].get(row["card"]) or cfg["cards"].get(row["card"].lstrip("0"))
        if card and card.get("scope"):
            return card["scope"]
        return "business?" if row["biz_hint"] else "personal"
    df["scope"] = df.apply(scope, axis=1)

    has_cards = (df["kind"] == "credit_card").any()
    ign = [s.upper() for s in cfg.get("ignore", [])]
    df["excluded"] = df["merchant"].str.upper().map(lambda m: any(s in m for s in ign))
    if has_cards:   # the card bill debited from the bank would count every purchase twice
        df.loc[df["category"] == "תשלום כרטיס אשראי", "excluded"] = True
    df["is_income"] = (df["category"] == "הכנסות") | ((df["kind"] == "bank") & (df["amount"] < 0))
    df["foreign"] = df["orig_currency"].str.upper().map(lambda c: c not in ("", "ILS", "NIS", "₪", "שח"))
    return df


def recurring(sp, last_month):
    """Fixed charges and subscriptions: same merchant in most months at a steady amount,
    standing orders, and yearly charges."""
    out = []
    regular = sp[~sp["txn_type"].isin(["installments", "credit"]) & (sp["amount"] > 0)]
    for key, g in regular.groupby("mkey"):
        by_month = g.groupby("month")["amount"].sum()
        per_month_count = g.groupby("month").size().mean()
        months = sorted(by_month.index)
        span = months_between(months[0], last_month) + 1
        standing = (g["txn_type"] == "standing_order").any()
        med = float(by_month.median())
        cv = float(by_month.std(ddof=0) / med) if med > 0 and len(by_month) > 1 else 0.0
        monthly = (len(months) >= 3 and len(months) >= 0.6 * span and cv <= 0.35
                   and per_month_count <= 2.2)
        yearly = False
        if not monthly and len(g) >= 2:
            ds = sorted(pd.to_datetime(g["date"]))
            amts = g.sort_values("date")["amount"].tolist()
            for i in range(1, len(ds)):
                gap = (ds[i] - ds[i - 1]).days
                if 330 <= gap <= 400 and abs(amts[i] - amts[i - 1]) <= 0.2 * max(amts[i - 1], 1):
                    yearly = True
        if not (monthly or yearly or (standing and len(months) >= 2)):
            continue
        last = g.sort_values("date").iloc[-1]
        earlier = by_month[by_month.index < months[-1]]
        change = None
        if len(earlier) >= 2:
            base = float(earlier.median())
            now = float(by_month.iloc[-1])
            if base > 0 and now - base >= 3 and (now - base) / base >= 0.05:
                change = {"from": round(base, 2), "to": round(now, 2), "pct": round(100 * (now - base) / base)}
        active = months_between(months[-1], last_month) <= (13 if yearly and not monthly else 1)
        out.append({
            "merchant": last["merchant"], "key": key, "category": last["category"],
            "nature": last["nature"], "scope": last["scope"], "card": last["card"],
            "frequency": "yearly" if yearly and not monthly else "monthly",
            "typical": round(med, 2), "last_amount": round(float(last["amount"]), 2),
            "monthly_equiv": round(med / (12 if yearly and not monthly else 1), 2),
            "annual_cost": round(med * (1 if yearly and not monthly else 12), 2),
            "months_seen": len(months), "first": g["date"].min(), "last": g["date"].max(),
            "standing_order": bool(standing), "active": bool(active), "price_change": change,
            "foreign": bool(g["foreign"].any()),
        })
    out.sort(key=lambda r: -r["annual_cost"])
    return out


def installments(df, last_month):
    plans, schedule = [], defaultdict(float)
    inst = df[df["txn_type"].isin(["installments", "credit"]) & df["installment_total"].notna()
              & (df["amount"] > 0)]
    for (key, card, total), g in inst.groupby(["mkey", "card", "installment_total"]):
        g = g.sort_values(["date", "installment_no"])
        # a merchant can have several plans with the same length: split by purchase date
        for d, p in g.groupby("date"):
            latest = p.sort_values("installment_no").iloc[-1]
            no, tot = int(latest["installment_no"]), int(total)
            remaining = max(tot - no, 0)
            amt = float(latest["amount"])
            start = latest["month"]
            plans.append({
                "merchant": latest["merchant"], "category": latest["category"], "scope": latest["scope"],
                "card": card, "purchase_date": d, "monthly": round(amt, 2), "paid_no": no, "total_no": tot,
                "remaining_payments": remaining, "remaining_amount": round(remaining * amt, 2),
                "full_price": round(float(latest["orig_amount"]), 2) if pd.notna(latest["orig_amount"]) else round(amt * tot, 2),
                "ends": add_months(start, remaining),
            })
            for k in range(1, remaining + 1):
                m = add_months(start, k)
                if m > last_month:
                    schedule[m] += amt
    plans.sort(key=lambda p: -p["remaining_amount"])
    return plans, {m: round(v, 2) for m, v in sorted(schedule.items())}


def anomalies(sp, rec, months, last_month):
    out = []
    pos = sp[sp["amount"] > 0]
    rec_keys = {r["key"] for r in rec}
    # 1. big one-off purchases
    if len(pos) >= 20:
        p95 = float(pos["amount"].quantile(0.95))
        thr = max(500.0, p95)
        # merchants seen in half the months or more are habits, covered by category spikes
        habitual = pos.groupby("mkey")["month"].nunique()
        habitual = set(habitual[habitual >= max(3, 0.5 * len(months))].index)
        big = pos[(pos["amount"] >= thr) & ~pos["mkey"].isin(rec_keys | habitual)
                  & ~pos["txn_type"].isin(["installments", "credit"])]
        for _, r in big.sort_values("amount", ascending=False).head(15).iterrows():
            out.append({"type": "big_purchase", "severity": "high" if r["amount"] >= 2 * thr else "medium",
                        "month": r["month"], "date": r["date"], "merchant": r["merchant"],
                        "amount": round(r["amount"], 2), "category": r["category"], "scope": r["scope"],
                        "detail": f"רכישה חד-פעמית גדולה (מעל {r0(thr)} ₪ - האחוזון ה-95 של ההוצאות שלך)"})
    # 2. category spikes in the last 3 months vs that category's own median
    if len(months) >= 4:
        pv = sp.pivot_table(index="month", columns="category", values="amount", aggfunc="sum").fillna(0)
        for m in months[-3:]:
            for cat in pv.columns:
                if CATEGORIES.get(cat) == "excluded":
                    continue
                others = pv.loc[[x for x in months if x != m and x in pv.index], cat]
                base = float(others.median())
                now = float(pv.loc[m, cat])
                if now - base >= 300 and now >= 1.4 * max(base, 1):
                    out.append({"type": "category_spike", "severity": "high" if now >= 2 * max(base, 1) else "medium",
                                "month": m, "category": cat, "amount": round(now, 2), "baseline": round(base, 2),
                                "detail": f"{cat}: {r0(now)} ₪ בחודש {m} מול חציון {r0(base)} ₪ בחודשים האחרים"})
    # 3. new merchants in the last month with a meaningful amount
    if len(months) >= 3:
        first_seen = pos.groupby("mkey")["month"].min()
        new = pos[(pos["month"] == last_month) & pos["mkey"].map(lambda k: first_seen.get(k) == last_month)]
        for key, g in new.groupby("mkey"):
            tot = float(g["amount"].sum())
            if tot >= 200:
                out.append({"type": "new_merchant", "severity": "low", "month": last_month,
                            "merchant": g["merchant"].iloc[0], "amount": round(tot, 2),
                            "category": g["category"].iloc[0], "detail": "בית עסק חדש החודש"})
    # 4. possible double charges: same merchant, same amount, same card, within 2 days
    reg = pos[pos["txn_type"].isin(["regular", "immediate", "standing_order"]) & (pos["amount"] >= 20)]
    for (key, amt, card), g in reg.groupby(["mkey", "amount", "card"]):
        if len(g) < 2:
            continue
        ds = sorted(pd.to_datetime(g["date"]))
        for i in range(1, len(ds)):
            if (ds[i] - ds[i - 1]).days <= 2:
                out.append({"type": "possible_duplicate", "severity": "medium", "month": g["month"].iloc[0],
                            "date": ds[i].date().isoformat(), "merchant": g["merchant"].iloc[0],
                            "amount": round(float(amt), 2), "card": card,
                            "detail": "אותו סכום באותו בית עסק פעמיים בתוך יומיים - לבדוק שזה לא חיוב כפול"})
                break
    # 5. subscription price increases
    for r in rec:
        if r["price_change"] and r["active"]:
            pc = r["price_change"]
            out.append({"type": "price_increase", "severity": "medium", "merchant": r["merchant"],
                        "amount": pc["to"], "baseline": pc["from"], "category": r["category"],
                        "detail": f"חיוב קבוע עלה מ-{pc['from']} ל-{pc['to']} ₪ (+{pc['pct']}%)"})
    # a spike that is just one big purchase already listed says the same thing twice
    big_cm = {(x["category"], x["month"]) for x in out if x["type"] == "big_purchase"}
    out = [x for x in out if not (x["type"] == "category_spike" and (x["category"], x["month"]) in big_cm)]
    order = {"high": 0, "medium": 1, "low": 2}
    out.sort(key=lambda a: (order[a["severity"]], -a.get("amount", 0)))
    return out


def baseline_budget(pv, months):
    """Per category: average, median, last month, and a suggested starting target.
    fixed/essential/business keep their median; lifestyle starts 15% under it. A category that is
    empty in most months (furniture, flights, clothes) is irregular: its monthly figure is the
    average, i.e. what to set aside each month so the occasional big bill is not a surprise."""
    rows = []
    n = len(months)
    for cat in pv.columns:
        nature = CATEGORIES.get(cat, "lifestyle")
        if nature == "excluded":
            continue
        s = pv[cat]
        med, avg = float(s.median()), float(s.sum() / n)
        irregular = (s > 0).sum() < 0.6 * n
        base = avg if irregular else med
        target = base * (0.85 if nature == "lifestyle" else 1.0)
        rows.append({"category": cat, "nature": nature, "irregular": bool(irregular),
                     "avg": round(avg), "median": round(med),
                     "last": round(float(s.iloc[-1])), "max": round(float(s.max())),
                     "suggested_target": int(round(target / 50.0) * 50) if nature in ("lifestyle", "essential")
                     and target >= 50 else round(target)})
    rows.sort(key=lambda r: -r["avg"])
    return rows


def forecast(sp, rec, schedule, full, k=3):
    """Next k billing months = active monthly fixed charges + yearly charges falling due +
    installments already committed + the median of everything else (the variable part)."""
    monthly_keys = {r["key"] for r in rec if r["frequency"] == "monthly"}
    rec_rows = sp["mkey"].isin(monthly_keys)
    inst_rows = sp["txn_type"].isin(["installments", "credit"])
    variable = sp[~rec_rows & ~inst_rows].groupby("month")["amount"].sum().reindex(full, fill_value=0)
    var_med = float(variable.median()) if len(variable) else 0.0
    fixed = sum(r["typical"] for r in rec if r["active"] and r["frequency"] == "monthly")
    out = []
    for i in range(1, k + 1):
        m = add_months(full[-1], i)
        yearly = [r for r in rec if r["frequency"] == "yearly" and r["active"]
                  and add_months(r["last"][:7], 12) == m]
        y_sum = sum(r["typical"] for r in yearly)
        inst = schedule.get(m, 0.0)
        out.append({"month": m, "fixed": round(fixed), "yearly_due": round(y_sum),
                    "yearly_items": [r["merchant"] for r in yearly], "installments": round(inst),
                    "variable_typical": round(var_med), "total": round(fixed + y_sum + inst + var_med)})
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("transactions")
    ap.add_argument("--config", default="finance/config.json")
    ap.add_argument("--out", default="finance/out")
    a = ap.parse_args()
    cfg = load_config(a.config)
    raw = pd.read_csv(a.transactions, dtype={"card": str}, encoding="utf-8-sig")
    if raw.empty:
        sys.exit("no transactions")
    df = prepare(raw, cfg)
    os.makedirs(a.out, exist_ok=True)

    sp = df[~df["excluded"] & ~df["is_income"]].copy()   # spending, refunds netted in
    months = sorted(sp["month"].unique())
    last_month = months[-1]
    n = len(months)

    # partial months (export cut mid-cycle) skew averages; flag months with < 40% of median rows
    counts = sp.groupby("month").size()
    partial = [m for m in months if counts[m] < 0.4 * counts.median()] if n >= 3 else []

    pv = sp.pivot_table(index="month", columns="category", values="amount", aggfunc="sum").fillna(0).reindex(months, fill_value=0)
    by_scope = sp.pivot_table(index="month", columns="scope", values="amount", aggfunc="sum").fillna(0).reindex(months, fill_value=0)
    monthly_total = pv.sum(axis=1)
    full = [m for m in months if m not in partial] or months

    rec = recurring(sp, last_month)
    plans, schedule = installments(sp, full[-1])
    anom = anomalies(sp, rec, full, full[-1])
    budget = baseline_budget(pv.loc[full], full)
    fc = forecast(sp, rec, schedule, full)

    # top merchants
    tm = sp[sp["amount"] > 0].groupby("mkey").agg(merchant=("merchant", "first"), total=("amount", "sum"),
                                                  count=("amount", "size"), category=("category", "first"),
                                                  scope=("scope", "first")).sort_values("total", ascending=False)
    top_merchants = [{"merchant": r.merchant, "total": round(r.total), "count": int(r.count),
                      "monthly_avg": round(r.total / n), "category": r.category, "scope": r.scope}
                     for r in tm.head(30).itertuples()]

    uncat = (sp[sp["category"] == UNCAT].groupby("mkey")
             .agg(merchant=("merchant", "first"), total=("amount", "sum"), count=("amount", "size"),
                  issuer_category=("issuer_category", "first"))
             .sort_values("total", ascending=False))
    uncategorized = [{"merchant": r.merchant, "total": round(r.total), "count": int(r.count),
                      "issuer_category": r.issuer_category} for r in uncat.head(40).itertuples()]

    suggested_biz = (sp[sp["scope"] == "business?"].groupby("mkey")
                     .agg(merchant=("merchant", "first"), total=("amount", "sum"), category=("category", "first"),
                          cards=("card", lambda s: ",".join(sorted(set(s)))))
                     .sort_values("total", ascending=False))
    biz = sp[sp["scope"].isin(["business", "business?"])]
    biz_dom = biz[~biz["foreign"] & (biz["amount"] > 0)]
    business = {
        "confirmed_total": round(float(sp.loc[sp["scope"] == "business", "amount"].sum())),
        "suggested_total": round(float(sp.loc[sp["scope"] == "business?", "amount"].sum())),
        "monthly_avg": round(float(biz["amount"].sum()) / n),
        "by_category": {k: round(v) for k, v in biz.groupby("category")["amount"].sum().sort_values(ascending=False).items()},
        "to_confirm": [{"merchant": r.merchant, "total": round(r.total), "category": r.category, "cards": r.cards}
                       for r in suggested_biz.head(40).itertuples()],
        # upper bound: only valid with a tax invoice, for an osek murshe, and for the recognised share
        "input_vat_upper_bound_domestic": round(float(biz_dom["amount"].sum()) * VAT_RATE / (1 + VAT_RATE)),
        "foreign_total": round(float(biz[biz["foreign"]]["amount"].sum())),
    }

    fees = sp[sp["category"] == "עמלות, ריבית ודמי כרטיס"]
    foreign = sp[sp["foreign"] & (sp["amount"] > 0)]
    subs = [r for r in rec if r["active"] and r["category"] in ("מנויים דיגיטליים", "תוכנה וכלי AI לעסק")]
    lifestyle_avg = {b["category"]: b["avg"] for b in budget if b["nature"] == "lifestyle"}
    income = df[df["is_income"] & ~df["excluded"]]
    opportunities = {
        "active_subscriptions_monthly": round(sum(r["monthly_equiv"] for r in subs)),
        "active_subscriptions_count": len(subs),
        "fixed_monthly": round(sum(r["monthly_equiv"] for r in rec if r["active"])),
        "fees_total": round(float(fees["amount"].sum())), "fees_monthly": round(float(fees["amount"].sum()) / n),
        "foreign_total": round(float(foreign["amount"].sum())), "foreign_count": int(len(foreign)),
        "eating_out_monthly": round(lifestyle_avg.get("מסעדות ובתי קפה", 0) + lifestyle_avg.get("משלוחי אוכל", 0)),
        "lifestyle_monthly": round(sum(lifestyle_avg.values())),
        "lifestyle_target_saving_monthly": round(sum((b["avg"] if b["irregular"] else b["median"]) - b["suggested_target"]
                                                     for b in budget if b["nature"] == "lifestyle")),
        "categories_with_3plus_subscriptions": sorted({r["category"] for r in subs if sum(1 for x in subs if x["category"] == r["category"]) >= 3}),
    }

    result = {
        "generated": dt.datetime.now().isoformat(timespec="seconds"),
        "period": {"first": df["date"].min(), "last": df["date"].max(), "months": months,
                   "partial_months": partial, "n_months": n},
        "sources": sorted(df["source_file"].unique().tolist()),
        "cards": {c: {"total": round(float(g["amount"].sum())), "count": int(len(g)),
                      "issuer": g["issuer"].iloc[0],
                      "label": (cfg["cards"].get(c) or {}).get("label", ""),
                      "scope": (cfg["cards"].get(c) or {}).get("scope", "")}
                  for c, g in sp.groupby("card")},
        "totals": {
            "spend": round(float(sp["amount"].sum())),
            "monthly_avg": round(float(monthly_total.loc[full].mean())),
            "monthly_median": round(float(monthly_total.loc[full].median())),
            "last_full_month": full[-1],
            "last_month": round(float(monthly_total[full[-1]])),
            "refunds": round(float(-sp.loc[sp["amount"] < 0, "amount"].sum())),
            "income_seen": round(float(-income["amount"].sum())),
            "excluded": round(float(df.loc[df["excluded"], "amount"].sum())),
        },
        "monthly": [{"month": m, "total": round(float(monthly_total[m])),
                     **{s: round(float(by_scope.loc[m, s])) for s in by_scope.columns}} for m in months],
        "category_month": {cat: [round(float(v)) for v in pv[cat]] for cat in pv.columns},
        "budget_baseline": budget,
        "recurring": rec,
        "installments": {"plans": plans, "future_by_month": schedule,
                         "remaining_total": round(sum(p["remaining_amount"] for p in plans))},
        "forecast": fc,
        "anomalies": anom,
        "top_merchants": top_merchants,
        "uncategorized": uncategorized,
        "uncategorized_share": round(100 * float(sp.loc[sp["category"] == UNCAT, "amount"].sum()) / max(float(sp["amount"].sum()), 1), 1),
        "business": business,
        "opportunities": opportunities,
        "user_budgets": cfg.get("budgets", {}),
        "user_income": cfg.get("income", {}),
    }
    with open(os.path.join(a.out, "analysis.json"), "w", encoding="utf-8") as fh:
        json.dump(result, fh, ensure_ascii=False, indent=1, default=str)

    cols = ["date", "charge_date", "month", "merchant", "amount", "category", "scope", "nature", "txn_type",
            "installment_no", "installment_total", "orig_amount", "orig_currency", "card", "issuer",
            "issuer_category", "notes", "excluded", "cat_source", "source_file"]
    df[cols].to_csv(os.path.join(a.out, "categorized.csv"), index=False, encoding="utf-8-sig")
    write_xlsx(os.path.join(a.out, "report.xlsx"), df[cols], result, pv)

    t = result["totals"]
    print(f"{n} months ({months[0]}..{last_month}), {len(sp)} spending rows, "
          f"avg {t['monthly_avg']:,} ₪/month, last {t['last_month']:,} ₪")
    print(f"fixed/recurring: {len(rec)} ({opportunities['fixed_monthly']:,} ₪/month active), "
          f"installment plans: {len(plans)} ({result['installments']['remaining_total']:,} ₪ still to pay), "
          f"anomalies: {len(anom)}, uncategorised: {result['uncategorized_share']}%")
    if partial:
        print(f"partial months (left out of averages): {', '.join(partial)}")
    print(f"-> {a.out}/analysis.json, categorized.csv, report.xlsx")


def write_xlsx(path, tx, res, pv):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    he = {"date": "תאריך", "charge_date": "תאריך חיוב", "month": "חודש", "merchant": "בית עסק",
          "amount": "סכום ₪", "category": "קטגוריה", "scope": "עסקי/פרטי", "nature": "סוג",
          "txn_type": "סוג עסקה", "installment_no": "תשלום", "installment_total": "מתוך",
          "orig_amount": "סכום מקורי", "orig_currency": "מטבע", "card": "כרטיס", "issuer": "חברה",
          "issuer_category": "קטגוריית החברה", "notes": "הערות", "excluded": "לא נספר",
          "cat_source": "מקור סיווג", "source_file": "קובץ"}
    sheets = {}
    sheets["סיכום חודשי"] = pd.DataFrame(res["monthly"]).rename(columns={
        "month": "חודש", "total": "סה\"כ", "personal": "פרטי", "business": "עסקי", "business?": "עסקי (לאישור)"})
    cm = pv.T.copy()
    cm["ממוצע"] = cm.mean(axis=1)
    sheets["קטגוריות לפי חודש"] = cm.sort_values("ממוצע", ascending=False).round(0).reset_index().rename(columns={"category": "קטגוריה"})
    sheets["תקציב בסיס"] = pd.DataFrame(res["budget_baseline"]).rename(columns={
        "category": "קטגוריה", "nature": "סוג", "avg": "ממוצע", "median": "חציון", "last": "חודש אחרון",
        "max": "מקסימום", "suggested_target": "יעד מוצע", "irregular": "לא סדיר"})
    sheets["קבועות ומנויים"] = pd.DataFrame(res["recurring"]).drop(columns=["key"], errors="ignore")
    sheets["תשלומים עתידיים"] = pd.DataFrame(res["installments"]["plans"])
    sheets["חריגות"] = pd.DataFrame(res["anomalies"])
    sheets["לאישור כעסקי"] = pd.DataFrame(res["business"]["to_confirm"])
    sheets["בתי עסק מובילים"] = pd.DataFrame(res["top_merchants"])
    sheets["לא מסווג"] = pd.DataFrame(res["uncategorized"])
    sheets["כל העסקאות"] = tx.rename(columns=he)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, d in sheets.items():
            (d if not d.empty else pd.DataFrame({"": ["אין נתונים"]})).to_excel(xw, sheet_name=name[:31], index=False)
        for ws in xw.book.worksheets:
            ws.sheet_view.rightToLeft = True
            ws.freeze_panes = "A2"
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="2A78D6")
                c.alignment = Alignment(horizontal="center", wrap_text=True)
            for i, col in enumerate(ws.columns, 1):
                width = max((len(str(c.value)) for c in list(col)[:200] if c.value is not None), default=8)
                ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 9), 42)
                for c in list(col)[1:]:
                    if isinstance(c.value, float):
                        c.number_format = "#,##0.00" if abs(c.value) < 100 else "#,##0"


if __name__ == "__main__":
    main()
